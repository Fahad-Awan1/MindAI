import os
import base64
import uuid
from typing import Optional
from PyPDF2 import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from io import BytesIO
from config import get_settings, BASE_DIR, resolve_path
from services import file_store

settings = get_settings()


def ensure_upload_dir():
    """Ensure upload directory exists."""
    os.makedirs(settings.resolved_upload_dir, exist_ok=True)


def resolve_stored_path(file_path: str) -> str:
    """
    Resolve a stored file path to an absolute path.

    Paths are stored relative to the backend directory (e.g. "./uploads/x.pdf")
    so the project keeps working if it is moved or the server is started from
    a different working directory.
    """
    return resolve_path(file_path)


def extract_text_from_pdf(file_data: bytes) -> str:
    """Extract text from PDF file."""
    try:
        pdf_reader = PdfReader(BytesIO(file_data))
        pages = []
        for page in pdf_reader.pages:
            try:
                # extract_text() can return None for pages without a text layer
                pages.append(page.extract_text() or "")
            except Exception as e:
                print(f"Error extracting text from PDF page: {e}")
        return "\n".join(pages).strip()
    except Exception as e:
        print(f"Error extracting PDF text: {e}")
        return ""


def extract_text_from_docx(file_data: bytes) -> str:
    """Extract text from DOCX file (paragraphs and tables)."""
    try:
        doc = DocxDocument(BytesIO(file_data))
        parts = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()
    except Exception as e:
        print(f"Error extracting DOCX text: {e}")
        return ""


def _extract_shape_text(shape) -> list[str]:
    """Recursively extract text from a PPTX shape (handles groups and tables)."""
    texts = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub_shape in shape.shapes:
            texts.extend(_extract_shape_text(sub_shape))
    elif getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                texts.append(" | ".join(cells))
    elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        if shape.text_frame.text.strip():
            texts.append(shape.text_frame.text)
    return texts


def extract_text_from_pptx(file_data: bytes) -> str:
    """Extract text from PPTX file."""
    try:
        prs = Presentation(BytesIO(file_data))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                parts.extend(_extract_shape_text(shape))
        return "\n".join(parts).strip()
    except Exception as e:
        print(f"Error extracting PPTX text: {e}")
        return ""


def process_document(file_content: bytes, file_type: str) -> tuple[str, Optional[str]]:
    """
    Process document content based on file type.

    Flow:
    1. Check file type (text, pdf, docx, etc.).
    2. For text/md: Decode directly.
    3. For binary (pdf, docx, pptx):
       - Extract text using specific libraries (PyPDF2, python-docx, python-pptx).
       - Save the binary file to disk (`save_binary_file`).
    4. Return tuple of (extracted_text, file_path).
    """
    if file_type in ['text', 'markdown']:
        # Text files - decode directly (utf-8-sig strips a BOM if present)
        try:
            return file_content.decode('utf-8-sig'), None
        except UnicodeDecodeError:
            # Fall back to latin-1, which can decode any byte sequence
            return file_content.decode('latin-1'), None

    extractors = {
        'pdf': extract_text_from_pdf,
        'docx': extract_text_from_docx,
        'pptx': extract_text_from_pptx,
    }
    extractor = extractors.get(file_type)
    if extractor is None:
        return "", None

    text_content = extractor(file_content)
    file_path = save_binary_file(file_content, file_type)
    return text_content, file_path


def save_binary_file(file_content: bytes, file_type: str) -> str:
    """Save binary file to disk and return its path (relative to the backend dir)."""
    ensure_upload_dir()
    filename = f"{uuid.uuid4().hex}.{file_type}"
    abs_path = os.path.join(settings.resolved_upload_dir, filename)

    with open(abs_path, 'wb') as f:
        f.write(file_content)

    stored = abs_path
    try:
        rel = os.path.relpath(abs_path, BASE_DIR)
        if not rel.startswith(".."):
            stored = "./" + rel.replace(os.sep, "/")
    except ValueError:
        # Different drive on Windows - fall back to absolute path
        pass
    # Also in the database: the disk may be wiped when the server restarts (cloud hosting)
    file_store.save(stored, file_content)
    return stored


def load_binary_file(file_path: str) -> Optional[bytes]:
    """Load a stored file: from disk, or from the database copy (written back to disk) when the disk lost it."""
    absolute = resolve_stored_path(file_path)
    try:
        with open(absolute, 'rb') as f:
            return f.read()
    except FileNotFoundError:
        pass
    except Exception as e:
        print(f"Error loading file: {e}")
        return None
    data = file_store.load(file_path)
    if data is not None:
        try:
            os.makedirs(os.path.dirname(absolute), exist_ok=True)
            with open(absolute, 'wb') as f:
                f.write(data)
        except OSError as e:
            print(f"Could not restore {file_path} to disk: {e}")
    return data


def get_file_base64(file_path: str) -> Optional[str]:
    """Get base64 encoded file data."""
    file_data = load_binary_file(file_path)
    if file_data:
        return base64.b64encode(file_data).decode('utf-8')
    return None


def delete_file(file_path: str):
    """Delete a stored file (disk and database copy)."""
    try:
        abs_path = resolve_stored_path(file_path)
        if os.path.exists(abs_path):
            os.remove(abs_path)
    except Exception as e:
        print(f"Error deleting file: {e}")
    try:
        file_store.delete(file_path)
    except Exception as e:
        print(f"Error deleting stored copy of {file_path}: {e}")


# ---------------------------------------------------------------------------
# Segments: the document split into labelled parts (pages, slides, sections)
# so citations can point to where a passage came from.
# ---------------------------------------------------------------------------

Segment = tuple  # (label: Optional[str], text: str)


def _pdf_segments(file_data: bytes) -> list:
    reader = PdfReader(BytesIO(file_data))
    segments = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception as e:
            print(f"Error extracting text from PDF page {i}: {e}")
            text = ""
        if text:
            segments.append((f"Page {i}", text))
    return segments


def _pptx_segments(file_data: bytes) -> list:
    prs = Presentation(BytesIO(file_data))
    segments = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            parts.extend(_extract_shape_text(shape))
        text = "\n".join(parts).strip()
        if text:
            segments.append((f"Slide {i}", text))
    return segments


def _docx_segments(file_data: bytes) -> list:
    doc = DocxDocument(BytesIO(file_data))
    segments = []
    label, lines = None, []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "") if paragraph.style is not None else ""
        if style.startswith("Heading") or style == "Title":
            if lines:
                segments.append((label, "\n".join(lines)))
            label, lines = text[:80], [text]
        else:
            lines.append(text)
    if lines:
        segments.append((label, "\n".join(lines)))
    table_rows = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_rows.append(" | ".join(cells))
    if table_rows:
        segments.append(("Tables", "\n".join(table_rows)))
    return segments


def _text_segments(text: str) -> list:
    """Split Markdown/plain text at headings ('# ...')."""
    segments = []
    label, lines = None, []
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            if any(l.strip() for l in lines):
                segments.append((label, "\n".join(lines).strip()))
            label, lines = stripped.lstrip("#").strip()[:80] or None, [line]
        else:
            lines.append(line)
    if any(l.strip() for l in lines):
        segments.append((label, "\n".join(lines).strip()))
    return segments


def extract_segments(file_content: bytes, file_type: str, fallback_text: str = "") -> list:
    """Labelled segments of a document; falls back to `fallback_text` as a single segment."""
    try:
        if file_type == "pdf":
            segments = _pdf_segments(file_content)
        elif file_type == "pptx":
            segments = _pptx_segments(file_content)
        elif file_type == "docx":
            segments = _docx_segments(file_content)
        else:
            segments = _text_segments(file_content.decode("utf-8-sig", errors="replace"))
    except Exception as e:
        print(f"Error extracting segments: {e}")
        segments = []
    if not segments and fallback_text.strip():
        segments = [(None, fallback_text.strip())]
    return segments


def document_segments(document) -> list:
    """Segments for a stored Document (re-reads the original file for PDF/DOCX/PPTX)."""
    if document.file_path and document.type in ("pdf", "docx", "pptx"):
        data = load_binary_file(document.file_path)
        if data:
            return extract_segments(data, document.type, document.content or "")
    return extract_segments((document.content or "").encode("utf-8"), "text", document.content or "")


def segments_to_text(segments: list) -> str:
    return "\n\n".join(text for _, text in segments).strip()
