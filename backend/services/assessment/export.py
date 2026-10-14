"""
Export quizzes and exams to other tools: Word (with a separate marking scheme), CSV, Anki,
Moodle (GIFT and XML), Kahoot and a Google Forms script. Version "B" shuffles the question
order (within sections) and the options, for a second paper that can't be copied from the first.
"""
import csv
import hashlib
import io
import json
import random
import re
import xml.etree.ElementTree as ET
from typing import Dict, List, Tuple

from services.assessment.grader import answer_text
from services.assessment.schema import TYPE_LABELS, upgrade_exam, upgrade_quiz

FORMATS = {
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "csv": ("text/csv", "csv"),
    "anki": ("text/tab-separated-values", "txt"),
    "gift": ("text/plain", "gift.txt"),
    "moodle": ("application/xml", "xml"),
    "kahoot": ("text/csv", "kahoot.csv"),
    "gforms": ("text/plain", "gs"),
}


def _rng(seed: str) -> random.Random:
    return random.Random(int(hashlib.sha1(seed.encode()).hexdigest()[:12], 16))


def _shuffle_options(q: Dict, rng: random.Random) -> Dict:
    q = dict(q)
    if q.get("type") in ("mcq", "multi_select") and q.get("options"):
        order = list(range(len(q["options"])))
        rng.shuffle(order)
        q["options"] = [q["options"][i] for i in order]
        if q.get("type") == "mcq" and isinstance(q.get("correct_answer_index"), int):
            q["correct_answer_index"] = order.index(q["correct_answer_index"])
        if q.get("type") == "multi_select":
            q["correct_indices"] = sorted(order.index(i) for i in q.get("correct_indices", []))
    return q


def load(study_set) -> Tuple[str, List[Dict], List[Dict], Dict]:
    """(kind, sections, questions, meta) for a quiz or exam set."""
    data = json.loads(study_set.data)
    if study_set.kind == "exam":
        exam = upgrade_exam(data)
        return "exam", exam["sections"], exam["questions"], {
            "title": exam.get("title") or study_set.title, "total_marks": exam.get("total_marks"),
            "duration": exam.get("duration_minutes"), "instructions": exam.get("instructions") or [],
            "header": exam.get("header")}
    questions = upgrade_quiz(data)
    return "quiz", [{"id": "A", "title": "", "instructions": "", "question_ids": [q["id"] for q in questions]}], questions, {
        "title": study_set.title, "total_marks": len(questions), "duration": None, "instructions": []}


def versioned(study_set, version: str) -> Tuple[str, List[Dict], List[Dict], Dict]:
    kind, sections, questions, meta = load(study_set)
    if version.upper() != "B":
        return kind, sections, questions, meta
    rng = _rng(study_set.id + "B")
    by_id = {q["id"]: _shuffle_options(q, rng) for q in questions}
    new_sections, ordered = [], []
    for s in sections:
        # Keep sub-parts of a question together and in order
        groups: Dict[str, List[str]] = {}
        for qid in s.get("question_ids", []):
            q = by_id.get(qid)
            if q:
                groups.setdefault(q.get("parent_id") or qid, []).append(qid)
        keys = list(groups)
        rng.shuffle(keys)
        ids = [qid for k in keys for qid in groups[k]]
        new_sections.append({**s, "question_ids": ids})
        ordered.extend(by_id[i] for i in ids)
    meta = {**meta, "title": f"{meta['title']} (Version B)"}
    return kind, new_sections, ordered, meta


def _numbered(sections: List[Dict], questions: List[Dict]) -> List[Tuple[Dict, Dict, str]]:
    """(section, question, number label) in paper order, renumbered (sub-parts as 3(a))."""
    by_id = {q["id"]: q for q in questions}
    out, n, parents, letters = [], 0, {}, {}
    for s in sections:
        for qid in s.get("question_ids", []):
            q = by_id.get(qid)
            if not q:
                continue
            p = q.get("parent_id")
            if p:
                if p not in parents:
                    n += 1
                    parents[p] = n
                letters[p] = letters.get(p, 0) + 1
                label = f"{parents[p]}({chr(96 + letters[p])})"
            else:
                n += 1
                label = str(n)
            out.append((s, q, label))
    return out


def _display_items(q: Dict) -> List[str]:
    """Ordering items in a shuffled (but stable) order for the paper."""
    items = list(q.get("items") or [])
    rng = _rng(q.get("id", "") + "order")
    shuffled = items[:]
    for _ in range(3):
        rng.shuffle(shuffled)
        if shuffled != items:
            break
    return shuffled


def _plain_answer(q: Dict) -> str:
    return answer_text(q) or q.get("answer_key", "")


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------

def to_docx(kind: str, sections: List[Dict], questions: List[Dict], meta: Dict, include_key: bool) -> bytes:
    from docx import Document as DocxDocument
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = DocxDocument()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    header = meta.get("header")
    if header and (header.get("institution") or header.get("fields")):
        _docx_header(doc, header, meta)
    else:
        _docx_title(doc, kind, questions, meta)

    current = None
    for section, q, label in _numbered(sections, questions):
        if kind == "exam" and section is not current:
            current = section
            doc.add_heading(f"Section {section['id']} · {section.get('title', '')}  [{section.get('marks', '')} marks]", level=1)
            if section.get("instructions"):
                doc.add_paragraph(section["instructions"]).italic = True
        if q.get("parent_stem") and label.endswith("(a)"):
            doc.add_paragraph(f"{label[:-3]}. {q['parent_stem']}")
        marks = f"  [{q.get('marks')}]" if kind == "exam" and q.get("marks") else ""
        qtype = q.get("type", "mcq")
        para = doc.add_paragraph()
        run = para.add_run(f"{label}. " if not q.get("parent_stem") else f"({label[-2:-1]}) ")
        run.bold = True
        para.add_run(q.get("question", "").replace("____", "__________") + marks)
        if qtype in ("mcq", "multi_select"):
            for i, opt in enumerate(q.get("options", [])):
                doc.add_paragraph(f"{chr(65 + i)})  {opt}")
            if qtype == "multi_select":
                doc.add_paragraph("(Select all that apply)").italic = True
        elif qtype == "true_false":
            doc.add_paragraph("True  /  False")
        elif qtype == "matching":
            table = doc.add_table(rows=0, cols=2)
            for left, right in zip(q.get("left", []), q.get("right", [])):
                cells = table.add_row().cells
                cells[0].text = f"____  {left}"
                cells[1].text = right
        elif qtype == "ordering":
            for item in _display_items(q):
                doc.add_paragraph(f"____  {item}")
        elif qtype == "numeric":
            doc.add_paragraph(f"Answer: ________________ {q.get('unit') or ''}")
        elif qtype in ("short_answer", "short", "long"):
            for _ in range(3 if qtype != "long" else 10):
                doc.add_paragraph("_" * 90)

    if include_key:
        doc.add_page_break()
        doc.add_heading("Answer key" if kind == "quiz" else "Marking scheme", level=0)
        for section, q, label in _numbered(sections, questions):
            para = doc.add_paragraph()
            para.add_run(f"{label}. ").bold = True
            if q.get("type") == "mcq":
                para.add_run(f"{chr(65 + q.get('correct_answer_index', 0))}) {_plain_answer(q)}")
            else:
                para.add_run(_plain_answer(q))
            for point in q.get("marking_scheme") or []:
                doc.add_paragraph(f"{point['point']}  ({point['marks']} mark{'s' if point['marks'] != 1 else ''})", style="List Bullet")
            if q.get("explanation") and kind == "quiz":
                note = doc.add_paragraph(q["explanation"])
                note.runs[0].italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _docx_title(doc, kind, questions, meta):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    title = doc.add_heading(meta["title"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info = []
    if meta.get("duration"):
        info.append(f"Time allowed: {meta['duration']} minutes")
    if kind == "exam":
        info.append(f"Maximum marks: {meta.get('total_marks')}")
    else:
        info.append(f"{len(questions)} questions")
    p = doc.add_paragraph(" · ".join(info))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Name: ______________________________        Date: ____________________")
    if meta.get("instructions"):
        doc.add_heading("Instructions", level=2)
        for line in meta["instructions"]:
            doc.add_paragraph(line, style="List Bullet")



def _docx_header(doc, header, meta):
    """University-style header: logo, institution lines, details table, instructions, exam name."""
    import base64
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    logo = header.get("logo")
    if logo and header.get("show_logo", True) and "," in logo:
        try:
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.add_run().add_picture(io.BytesIO(base64.b64decode(logo.split(",", 1)[1])), width=Inches(0.9))
        except Exception:
            pass
    for text, size in ((header.get("institution"), 18), (header.get("campus"), 14), (header.get("department"), 14)):
        if text:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(size)
    fields = []
    for f in header.get("fields") or []:
        value = f.get("value") or ""
        if not value and f.get("auto") == "duration" and meta.get("duration"):
            value = f"{meta['duration']} Minutes"
        if not value and f.get("auto") == "marks":
            value = str(meta.get("total_marks") or "")
        fields.append((f.get("label", ""), value))
    if fields:
        rows = (len(fields) + 1) // 2
        table = doc.add_table(rows=rows, cols=4)
        table.style = "Table Grid"
        for i, (label, value) in enumerate(fields):
            r, c = divmod(i, 2)
            cells = table.rows[r].cells
            cells[c * 2].text = label
            cells[c * 2].paragraphs[0].runs[0].bold = True
            cells[c * 2 + 1].text = value
    if meta.get("instructions"):
        p = doc.add_paragraph()
        p.add_run("Important Instructions / Guidelines:").bold = True
        for line in meta["instructions"]:
            doc.add_paragraph(line).runs[0].italic = True
    name = header.get("exam_name") or meta["title"]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(name)
    run.bold = True
    run.underline = True
    run.font.size = Pt(14)


# ---------------------------------------------------------------------------
# Text formats
# ---------------------------------------------------------------------------

def to_csv(sections, questions, meta) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["number", "type", "question", "options", "answer", "marks", "explanation", "topic", "difficulty", "bloom", "source"])
    for _, q, label in _numbered(sections, questions):
        if q.get("type") == "ordering":
            options = _display_items(q)
        elif q.get("type") == "matching":
            options = list(q.get("left") or []) + list(q.get("right") or [])
        else:
            options = q.get("options") or []
        src = q.get("source") or {}
        w.writerow([label, TYPE_LABELS.get(q.get("type"), q.get("type")), q.get("question", ""), " | ".join(options or []),
                    _plain_answer(q), q.get("marks", ""), q.get("explanation", ""), q.get("topic", ""), q.get("difficulty", ""),
                    q.get("bloom", ""), " · ".join(x for x in [src.get("document_title"), src.get("location")] if x)])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _html(text: str) -> str:
    return (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br>").replace("\t", " ")


def to_anki(sections, questions, meta) -> bytes:
    lines = ["#separator:tab", "#html:true", "#tags column:3"]
    deck_tag = re.sub(r"\W+", "_", meta["title"]).strip("_")[:40] or "MindAI"
    for _, q, _ in _numbered(sections, questions):
        front = _html(q.get("question", ""))
        if q.get("options") and q.get("type") in ("mcq", "multi_select"):
            front += "<br><br>" + "<br>".join(f"{chr(65 + i)}) {_html(o)}" for i, o in enumerate(q["options"]))
        if q.get("type") == "matching":
            front += "<br><br>" + "<br>".join(_html(x) for x in q.get("left", []))
        back = f"<b>{_html(_plain_answer(q))}</b>"
        if q.get("explanation"):
            back += "<br><br>" + _html(q["explanation"])
        tags = " ".join(t for t in [deck_tag, re.sub(r"\W+", "_", q.get("topic") or "")] if t)
        lines.append(f"{front}\t{back}\t{tags}")
    return ("\n".join(lines) + "\n").encode("utf-8")


def _gift_escape(text: str) -> str:
    return re.sub(r"([~=#{}:])", r"\\\1", str(text or "")).replace("\n", " ")


def to_gift(sections, questions, meta) -> bytes:
    out = [f"// {meta['title']} - exported from MindAI", ""]
    for _, q, label in _numbered(sections, questions):
        name = f"::Q{label} {_gift_escape((q.get('topic') or '')[:40])}::"
        stem = _gift_escape(q.get("question", ""))
        fb = f"####{_gift_escape(q['explanation'])}" if q.get("explanation") else ""
        t = q.get("type", "mcq")
        if t == "mcq":
            opts = " ".join(("=" if i == q.get("correct_answer_index") else "~") + _gift_escape(o) for i, o in enumerate(q.get("options", [])))
            out.append(f"{name}{stem} {{{opts} {fb}}}")
        elif t == "multi_select":
            right = q.get("correct_indices", [])
            wrong_n = max(1, len(q.get("options", [])) - len(right))
            pct_r, pct_w = round(100 / max(1, len(right)), 5), round(100 / wrong_n, 5)
            opts = " ".join((f"~%{pct_r}%" if i in right else f"~%-{pct_w}%") + _gift_escape(o) for i, o in enumerate(q.get("options", [])))
            out.append(f"{name}{stem} {{{opts} {fb}}}")
        elif t == "true_false":
            out.append(f"{name}{stem} {{{'T' if q.get('correct_bool') else 'F'}{fb}}}")
        elif t == "fill_blank" and len(q.get("blanks", [])) == 1:
            parts = q["question"].split("____")
            answers = " ".join("=" + _gift_escape(a) for a in q["blanks"][0]["accepted"])
            out.append(f"{name}{_gift_escape(parts[0])} {{{answers}}} {_gift_escape(parts[1] if len(parts) > 1 else '')}")
        elif t in ("fill_blank", "short_answer"):
            accepted = [a for b in q.get("blanks", []) for a in b["accepted"][:1]] if t == "fill_blank" else [q.get("model_answer", "")]
            out.append(f"{name}{stem} {{{' '.join('=' + _gift_escape(a) for a in accepted)} {fb}}}")
        elif t == "numeric":
            out.append(f"{name}{stem} {{#{q.get('answer')}:{q.get('tolerance', 0)} {fb}}}")
        elif t in ("matching", "ordering"):
            if t == "matching":
                pairs = [(q["left"][li], q["right"][ri]) for li, ri in q.get("pairs", [])]
            else:
                pairs = [(item, f"Position {i + 1}") for i, item in enumerate(q.get("items", []))]
            body = " ".join(f"={_gift_escape(a)} -> {_gift_escape(b)}" for a, b in pairs)
            out.append(f"{name}{stem} {{{body}}}")
        else:  # short / long written answers -> essay
            out.append(f"{name}{stem} {{}}")
        out.append("")
    return "\n".join(out).encode("utf-8")


def _cdata_text(parent, tag: str, text: str, fmt: str = "html"):
    el = ET.SubElement(parent, tag, {"format": fmt} if fmt else {})
    ET.SubElement(el, "text").text = text
    return el


def to_moodle_xml(sections, questions, meta) -> bytes:
    quiz = ET.Element("quiz")
    for _, q, label in _numbered(sections, questions):
        t = q.get("type", "mcq")
        mtype = {"mcq": "multichoice", "multi_select": "multichoice", "true_false": "truefalse", "numeric": "numerical",
                 "matching": "matching", "ordering": "matching", "fill_blank": "shortanswer", "short_answer": "shortanswer"}.get(t, "essay")
        el = ET.SubElement(quiz, "question", {"type": mtype})
        name = ET.SubElement(el, "name")
        ET.SubElement(name, "text").text = f"Q{label} {q.get('topic') or ''}".strip()
        _cdata_text(el, "questiontext", _html(q.get("question", "")))
        if q.get("explanation"):
            _cdata_text(el, "generalfeedback", _html(q["explanation"]))
        ET.SubElement(el, "defaultgrade").text = str(q.get("marks") or 1)
        if t in ("mcq", "multi_select"):
            right = [q.get("correct_answer_index")] if t == "mcq" else q.get("correct_indices", [])
            ET.SubElement(el, "single").text = "true" if t == "mcq" else "false"
            ET.SubElement(el, "shuffleanswers").text = "true"
            for i, opt in enumerate(q.get("options", [])):
                fraction = (100 / max(1, len(right))) if i in right else (0 if t == "mcq" else -100 / max(1, len(q["options"]) - len(right)))
                ans = ET.SubElement(el, "answer", {"fraction": f"{fraction:.5g}", "format": "html"})
                ET.SubElement(ans, "text").text = _html(opt)
        elif t == "true_false":
            for value in ("true", "false"):
                ans = ET.SubElement(el, "answer", {"fraction": "100" if (value == "true") == bool(q.get("correct_bool")) else "0"})
                ET.SubElement(ans, "text").text = value
        elif t == "numeric":
            ans = ET.SubElement(el, "answer", {"fraction": "100"})
            ET.SubElement(ans, "text").text = str(q.get("answer"))
            ET.SubElement(ans, "tolerance").text = str(q.get("tolerance", 0))
        elif t in ("matching", "ordering"):
            pairs = [(q["left"][li], q["right"][ri]) for li, ri in q.get("pairs", [])] if t == "matching" else \
                [(item, f"Position {i + 1}") for i, item in enumerate(q.get("items", []))]
            for a, b in pairs:
                sub = ET.SubElement(el, "subquestion", {"format": "html"})
                ET.SubElement(sub, "text").text = _html(a)
                ans = ET.SubElement(sub, "answer")
                ET.SubElement(ans, "text").text = b
        elif t in ("fill_blank", "short_answer"):
            accepted = [a for b in q.get("blanks", []) for a in b["accepted"]] if t == "fill_blank" else [q.get("model_answer", "")]
            for a in accepted[:6]:
                ans = ET.SubElement(el, "answer", {"fraction": "100"})
                ET.SubElement(ans, "text").text = a
        else:
            _cdata_text(el, "graderinfo", _html(q.get("answer_key", "")))
    ET.indent(quiz)
    return b'<?xml version="1.0" encoding="UTF-8"?>\n' + ET.tostring(quiz, encoding="utf-8")


KAHOOT_TIMES = [5, 10, 20, 30, 60, 90, 120, 240]


def to_kahoot(sections, questions, meta) -> bytes:
    """Kahoot's spreadsheet template columns (choice questions only, up to 4 answers)."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Question - max 120 characters", "Answer 1 - max 75 characters", "Answer 2 - max 75 characters",
                "Answer 3 - max 75 characters", "Answer 4 - max 75 characters",
                "Time limit (sec) – 5, 10, 20, 30, 60, 90, 120, or 240 secs", "Correct answer(s) - choose at least one"])
    for _, q, _ in _numbered(sections, questions):
        t = q.get("type")
        if t not in ("mcq", "true_false", "multi_select"):
            continue
        options = (q.get("options") or [])[:4]
        if t == "mcq":
            correct = [q.get("correct_answer_index", 0)]
        elif t == "true_false":
            correct = [0 if q.get("correct_bool") else 1]
        else:
            correct = [i for i in q.get("correct_indices", []) if i < 4]
        if not correct or max(correct) >= len(options):
            continue
        time_limit = 20 if len(q.get("question", "")) < 120 else 30
        w.writerow([q.get("question", "")[:120], *[o[:75] for o in options], *[""] * (4 - len(options)), time_limit,
                    ",".join(str(i + 1) for i in correct)])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_gforms(sections, questions, meta, kind: str) -> bytes:
    """A Google Apps Script that builds this as a self-marking Google Forms quiz (paste into script.google.com)."""
    items = []
    for _, q, label in _numbered(sections, questions):
        items.append({
            "label": label, "type": q.get("type", "mcq"), "question": q.get("question", ""), "options": q.get("options") or [],
            "correct": [q.get("correct_answer_index")] if q.get("type") == "mcq" else q.get("correct_indices") or [],
            "bool": q.get("correct_bool"), "answer": _plain_answer(q), "points": q.get("marks") or 1,
            "left": q.get("left") or [], "right": q.get("right") or [], "pairs": q.get("pairs") or [],
            "items": q.get("items") or [], "explanation": q.get("explanation", ""),
        })
    data = json.dumps({"title": meta["title"], "items": items}, ensure_ascii=False, indent=1)
    script = f"""/**
 * {meta['title']} - exported from MindAI.
 * How to use: open https://script.google.com, create a new project, paste this file, then press Run.
 * It creates a self-marking Google Forms quiz in your Google Drive and logs its link.
 */
const DATA = {data};

function createQuiz() {{
  const form = FormApp.create(DATA.title).setIsQuiz(true);
  DATA.items.forEach(function (q) {{
    const title = q.label + '. ' + q.question;
    if (q.type === 'mcq' || q.type === 'true_false') {{
      const item = form.addMultipleChoiceItem().setTitle(title).setPoints(q.points);
      const opts = q.type === 'true_false' ? ['True', 'False'] : q.options;
      const right = q.type === 'true_false' ? (q.bool ? 0 : 1) : q.correct[0];
      item.setChoices(opts.map(function (o, i) {{ return item.createChoice(o, i === right); }}));
    }} else if (q.type === 'multi_select') {{
      const item = form.addCheckboxItem().setTitle(title).setPoints(q.points);
      item.setChoices(q.options.map(function (o, i) {{ return item.createChoice(o, q.correct.indexOf(i) >= 0); }}));
    }} else if (q.type === 'matching') {{
      const grid = form.addGridItem().setTitle(title).setRows(q.left).setColumns(q.right);
    }} else if (q.type === 'long' || q.type === 'short') {{
      form.addParagraphTextItem().setTitle(title).setPoints(q.points);
    }} else {{
      form.addTextItem().setTitle(title).setPoints(q.points);
    }}
  }});
  Logger.log('Edit: ' + form.getEditUrl());
  Logger.log('Share: ' + form.getPublishedUrl());
}}
"""
    return script.encode("utf-8")


def export_set(study_set, fmt: str, version: str = "A", include_key: bool = True) -> Tuple[bytes, str, str]:
    """(content, media type, filename)."""
    if fmt not in FORMATS:
        raise ValueError(f"Unknown format '{fmt}'")
    kind, sections, questions, meta = versioned(study_set, version)
    if fmt == "docx":
        content = to_docx(kind, sections, questions, meta, include_key)
    elif fmt == "csv":
        content = to_csv(sections, questions, meta)
    elif fmt == "anki":
        content = to_anki(sections, questions, meta)
    elif fmt == "gift":
        content = to_gift(sections, questions, meta)
    elif fmt == "moodle":
        content = to_moodle_xml(sections, questions, meta)
    elif fmt == "kahoot":
        content = to_kahoot(sections, questions, meta)
    else:
        content = to_gforms(sections, questions, meta, kind)
    media, ext = FORMATS[fmt]
    base = re.sub(r"[^\w\- ]+", "", meta["title"]).strip().replace(" ", "_")[:60] or "mindai"
    return content, media, f"{base}.{ext}"
